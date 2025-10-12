import json
from pathlib import Path
from pptx import Presentation

OUTLINE_PATH = Path("/workspace/slides/zoho-creator-prep-outline.json")
OUTPUT_PATH = Path("/workspace/slides/zoho-creator-prep-slides.pptx")


def add_bullets(text_frame, bullets):
    if not bullets:
        return
    # initialize first line
    text_frame.text = bullets[0]
    # subsequent bullets as paragraphs
    for bullet in bullets[1:]:
        p = text_frame.add_paragraph()
        p.text = bullet
        p.level = 0


def main():
    with OUTLINE_PATH.open("r", encoding="utf-8") as f:
        data = json.load(f)

    prs = Presentation()

    for idx, slide_data in enumerate(data.get("slides", [])):
        title = slide_data.get("title", "")
        bullets = slide_data.get("bullets", [])
        notes = slide_data.get("notes", "")

        # First slide as Title slide if outline uses generic "Title"
        if idx == 0 and title.strip().lower() == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title layout
            # Title and subtitle from first two bullets if present
            if bullets:
                slide.shapes.title.text = bullets[0]
            if len(bullets) > 1:
                subtitle_placeholder = slide.placeholders[1]
                subtitle_placeholder.text = bullets[1]
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
            slide.shapes.title.text = title
            content_ph = slide.placeholders[1]
            tf = content_ph.text_frame
            tf.clear()
            add_bullets(tf, bullets)

        # Speaker notes
        if notes:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = notes

    prs.save(OUTPUT_PATH)
    print(str(OUTPUT_PATH))


if __name__ == "__main__":
    main()
